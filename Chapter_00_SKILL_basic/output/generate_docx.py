#!/usr/bin/env python3
"""Generate a DOCX version of the test plan from the markdown file.

Usage:
  python3 generate_docx.py

Output: testplan_vwo_login.docx
"""
from docx import Document
from docx.shared import Pt
import os

ROOT = os.path.dirname(__file__)
MD = os.path.join(ROOT, 'testplan_vwo_login.md')
OUT = os.path.join(ROOT, 'testplan_vwo_login.docx')

def md_to_docx(md_path, out_path):
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    with open(md_path, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.rstrip('\n')
            if not line.strip():
                doc.add_paragraph('')
                continue
            if line.startswith('# '):
                doc.add_heading(line[2:].strip(), level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:].strip(), level=3)
            elif line.startswith('- ') or line.startswith('+ '):
                p = doc.add_paragraph(line[2:].strip(), style='List Bullet')
            else:
                doc.add_paragraph(line)

    doc.save(out_path)
    print('DOCX generated at', out_path)

if __name__ == '__main__':
    if not os.path.exists(MD):
        print('Markdown file not found:', MD)
        raise SystemExit(1)
    md_to_docx(MD, OUT)
